"""
Statement of account files (7.17): Excel (openpyxl) and PDF (reportlab).

Both are generated server-side from the same build_statement() result, so a
downloaded file and an emailed one are always identical. Customer-facing:
no internal system names anywhere in the output.
"""
import io
import os
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from reportlab.lib import colors
from reportlab.lib.enums import TA_RIGHT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import (
    Image, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle,
)

from services.customer_statement import statement_period_label

BRAND = "0F6E56"          # tailwind bassani-600
BRAND_LIGHT = "E1F5EE"     # bassani-50
GREY = "6B7280"
RED = "B91C1C"

_HERE = os.path.dirname(os.path.abspath(__file__))
_LOGO_CANDIDATES = [
    os.path.join(_HERE, "..", "static", "logo.png"),                         # production (/app/static)
    os.path.join(_HERE, "..", "..", "frontend", "public", "logo.png"),       # local dev
]


def _logo_path():
    for p in _LOGO_CANDIDATES:
        if os.path.isfile(p):
            return p
    return None


def _d(s):
    if not s:
        return ""
    try:
        return datetime.strptime(str(s)[:10], "%Y-%m-%d").strftime("%d %b %Y")
    except ValueError:
        return str(s)


def _money(v):
    v = v or 0.0
    s = f"R{abs(v):,.2f}"
    return f"-{s}" if v < 0 else s


def _payment_reference(stmt):
    c = stmt.get("customer") or {}
    return c.get("ref") or c.get("name") or ""


_TOTAL_LABELS = ("Closing balance", "Total due", "Balance in credit")


def _summary_pairs(stmt):
    s = stmt.get("summary") or {}
    if stmt.get("mode") == "open":
        total = s.get("total_due", 0) or 0
        pairs = [("Overdue", s.get("overdue", 0))]
        if abs(s.get("unallocated") or 0) >= 0.005:
            pairs.append(("Unallocated payments and credits", s.get("unallocated", 0)))
        pairs.append(("Total due" if total >= 0 else "Balance in credit", total))
        return pairs
    pairs = [
        ("Opening balance", s.get("opening_balance", 0)),
        ("Invoices", s.get("invoiced", 0)),
        ("Payments received", -(s.get("payments", 0) or 0)),
        ("Credit notes", -(s.get("credit_notes", 0) or 0)),
    ]
    if s.get("adjustments"):
        pairs.append(("Adjustments", s.get("adjustments", 0)))
    pairs.append(("Closing balance", s.get("closing_balance", 0)))
    return pairs


def _ageing_cells(ageing):
    """Ageing columns plus, when the customer has payments or credit notes
    not yet matched to an invoice, an Unallocated column, so the strip adds
    up to the total shown at the end of it."""
    cells = [(b["label"], b["amount"]) for b in ageing["buckets"]]
    if abs(ageing.get("unallocated") or 0) >= 0.005:
        cells.append(("Unallocated credits", ageing["unallocated"]))
    total = ageing["total_due"]
    cells.append(("Total due" if total >= 0 else "Balance in credit", total))
    return cells


def _balance_note(stmt):
    total = (stmt.get("ageing") or {}).get("total_due", 0) or 0
    if total < -0.005:
        return f"Your account is in credit by {_money(-total)}."
    if total < 0.005:
        return "Your account is fully paid. Thank you."
    return (f"Amount due today: {_money(total)}. Please use "
            f"\"{_payment_reference(stmt)}\" as your payment reference.")


# ── Excel ─────────────────────────────────────────────────────────────────────

def build_statement_xlsx(stmt: dict) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Statement"
    company = stmt.get("company") or {}
    customer = stmt.get("customer") or {}
    is_open = stmt.get("mode") == "open"

    money_fmt = '"R"#,##0.00;[Red]-"R"#,##0.00'
    bold = Font(bold=True)
    muted = Font(color=GREY, size=9)
    head_fill = PatternFill("solid", fgColor=BRAND)
    head_font = Font(bold=True, color="FFFFFF")
    thin = Side(style="thin", color="D1D5DB")

    ws["A1"] = "Statement of Account" if not is_open else "Statement of Open Items"
    ws["A1"].font = Font(bold=True, size=16, color=BRAND)
    ws["A2"] = statement_period_label(stmt)
    ws["A2"].font = muted

    row = 4
    ws.cell(row=row, column=1, value="From").font = muted
    ws.cell(row=row, column=5, value="To").font = muted
    left = [company.get("name"), company.get("address"),
            f"VAT: {company['vat']}" if company.get("vat") else None,
            f"Reg: {company['company_registry']}" if company.get("company_registry") else None,
            company.get("email"), company.get("phone")]
    right = [customer.get("name"),
             ", ".join(x for x in [customer.get("street"), customer.get("city"), customer.get("zip")] if x) or None,
             f"VAT: {customer['vat']}" if customer.get("vat") else None,
             f"Account: {customer['ref']}" if customer.get("ref") else None]
    left = [x for x in left if x]
    right = [x for x in right if x]
    for i in range(max(len(left), len(right))):
        r = row + 1 + i
        if i < len(left):
            ws.cell(row=r, column=1, value=left[i]).font = bold if i == 0 else Font()
        if i < len(right):
            ws.cell(row=r, column=5, value=right[i]).font = bold if i == 0 else Font()
    row += max(len(left), len(right)) + 2

    ws.cell(row=row, column=1, value="Summary").font = Font(bold=True, color=BRAND)
    row += 1
    for label, value in _summary_pairs(stmt):
        ws.cell(row=row, column=1, value=label)
        c = ws.cell(row=row, column=3, value=round(value or 0, 2))
        c.number_format = money_fmt
        if label in _TOTAL_LABELS:
            ws.cell(row=row, column=1).font = bold
            c.font = bold
        row += 1
    row += 1

    ageing = stmt.get("ageing")
    if ageing:
        ws.cell(row=row, column=1, value="Ageing (as at " + _d(stmt.get("as_of")) + ")").font = Font(bold=True, color=BRAND)
        row += 1
        cells = _ageing_cells(ageing)
        for i, (label, amount) in enumerate(cells):
            last = i == len(cells) - 1
            ws.cell(row=row, column=1 + i, value=label).font = Font(bold=True, size=9) if last else muted
            c = ws.cell(row=row + 1, column=1 + i, value=amount)
            c.number_format = money_fmt
            if last:
                c.font = bold
        row += 3

    if is_open:
        headers = ["Date", "Number", "Type", "Reference", "Due date", "Days overdue", "Outstanding"]
    else:
        headers = ["Date", "Number", "Type", "Reference", "Due date", "Debit", "Credit", "Balance"]
    header_row = row
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=row, column=i, value=h)
        c.fill = head_fill
        c.font = head_font
        c.alignment = Alignment(horizontal="right" if i > 5 else "left")
    row += 1

    if not is_open and stmt.get("date_from"):
        ws.cell(row=row, column=1, value=_d(stmt["date_from"]))
        ws.cell(row=row, column=4, value="Opening balance").font = Font(italic=True)
        c = ws.cell(row=row, column=8, value=(stmt.get("summary") or {}).get("opening_balance", 0))
        c.number_format = money_fmt
        row += 1

    for r in stmt.get("rows", []):
        vals = [_d(r.get("date")), r.get("number"), r.get("type_label"), r.get("reference"),
                _d(r.get("due_date")) if r.get("kind") in ("invoice", "credit_note") else ""]
        if is_open:
            owed = (r.get("residual") or 0) > 0
            vals += [(r.get("days_overdue") or 0) if owed else None, r.get("residual")]
        else:
            vals += [r.get("debit") or None, r.get("credit") or None, r.get("running_balance")]
        money_cols = (7,) if is_open else (6, 7, 8)
        for i, v in enumerate(vals, start=1):
            c = ws.cell(row=row, column=i, value=v)
            c.border = Border(bottom=thin)
            if i in money_cols and v is not None:
                c.number_format = money_fmt
        if is_open and (r.get("days_overdue") or 0) > 0 and (r.get("residual") or 0) > 0:
            ws.cell(row=row, column=6).font = Font(color=RED, bold=True)
        row += 1

    if stmt.get("truncated"):
        row += 1
        ws.cell(row=row, column=1, value="This statement is too long to show in full. Please choose a shorter period.").font = Font(color=RED)

    row += 2
    ws.cell(row=row, column=1, value=_balance_note(stmt)).font = bold
    bank = company.get("bank") or {}
    if bank.get("account"):
        row += 1
        ws.cell(row=row, column=1, value=f"Banking details: {bank['account']}")
        if bank.get("holder"):
            row += 1
            ws.cell(row=row, column=1, value=f"Account holder: {bank['holder']}")
    row += 1
    ws.cell(row=row, column=1, value=f"Generated {_d(stmt.get('generated_at'))}").font = muted

    widths = [13, 20, 17, 30, 13, 15, 15, 16]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    ws.print_options.gridLines = False
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToHeight = 0

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ── PDF ───────────────────────────────────────────────────────────────────────

def build_statement_pdf(stmt: dict) -> bytes:
    company = stmt.get("company") or {}
    customer = stmt.get("customer") or {}
    is_open = stmt.get("mode") == "open"
    brand = colors.HexColor("#" + BRAND)
    grey = colors.HexColor("#" + GREY)

    styles = getSampleStyleSheet()
    base = ParagraphStyle("base", parent=styles["Normal"], fontName="Helvetica", fontSize=8.5, leading=11)
    small = ParagraphStyle("small", parent=base, fontSize=7.5, textColor=grey, leading=10)
    bold = ParagraphStyle("bold", parent=base, fontName="Helvetica-Bold")
    title = ParagraphStyle("title", parent=base, fontName="Helvetica-Bold", fontSize=16, leading=20, textColor=brand)
    right = ParagraphStyle("right", parent=base, alignment=TA_RIGHT)

    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=A4, leftMargin=14 * mm, rightMargin=14 * mm,
        topMargin=14 * mm, bottomMargin=16 * mm,
        title="Statement of Account", author=company.get("name") or "",
    )
    width = A4[0] - 28 * mm
    story = []

    # Letterhead: logo left, issuing company right.
    logo = _logo_path()
    logo_cell = ""
    if logo:
        try:
            img = Image(logo)
            ratio = img.imageHeight / float(img.imageWidth or 1)
            img.drawWidth = 52 * mm
            img.drawHeight = 52 * mm * ratio
            if img.drawHeight > 22 * mm:
                img.drawHeight = 22 * mm
                img.drawWidth = 22 * mm / ratio
            logo_cell = img
        except Exception:
            logo_cell = ""
    co_lines = [f"<b>{company.get('name', '')}</b>"]
    for x in [company.get("address"),
              f"VAT {company['vat']}" if company.get("vat") else None,
              f"Reg {company['company_registry']}" if company.get("company_registry") else None,
              " | ".join(y for y in [company.get("email"), company.get("phone")] if y) or None]:
        if x:
            co_lines.append(x)
    head = Table([[logo_cell, Paragraph("<br/>".join(co_lines), right)]],
                 colWidths=[width * 0.45, width * 0.55])
    head.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"),
                              ("LEFTPADDING", (0, 0), (-1, -1), 0),
                              ("RIGHTPADDING", (0, 0), (-1, -1), 0)]))
    story += [head, Spacer(1, 6 * mm)]

    story.append(Paragraph("Statement of Account" if not is_open else "Statement of Open Items", title))
    story.append(Paragraph(statement_period_label(stmt), small))
    story.append(Spacer(1, 4 * mm))

    cust_lines = [f"<b>{customer.get('name', '')}</b>"]
    addr = ", ".join(x for x in [customer.get("street"), customer.get("street2"), customer.get("city"), customer.get("zip")] if x)
    for x in [addr or None, f"VAT {customer['vat']}" if customer.get("vat") else None,
              f"Account {customer['ref']}" if customer.get("ref") else None]:
        if x:
            cust_lines.append(x)
    summary_rows = [[Paragraph(label, bold if label in _TOTAL_LABELS else base),
                     Paragraph(_money(value), ParagraphStyle("r", parent=bold if label in _TOTAL_LABELS else base, alignment=TA_RIGHT))]
                    for label, value in _summary_pairs(stmt)]
    summary_tbl = Table(summary_rows, colWidths=[width * 0.28, width * 0.17])
    summary_tbl.setStyle(TableStyle([
        ("LINEABOVE", (0, -1), (-1, -1), 0.6, brand),
        ("TOPPADDING", (0, 0), (-1, -1), 1.5), ("BOTTOMPADDING", (0, 0), (-1, -1), 1.5),
    ]))
    info = Table([[Paragraph("<font color='#6B7280' size='7.5'>STATEMENT FOR</font><br/>" + "<br/>".join(cust_lines), base),
                   summary_tbl]],
                 colWidths=[width * 0.55, width * 0.45])
    info.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"),
                              ("BACKGROUND", (0, 0), (0, 0), colors.HexColor("#" + BRAND_LIGHT)),
                              ("LEFTPADDING", (0, 0), (0, 0), 8), ("TOPPADDING", (0, 0), (0, 0), 6),
                              ("BOTTOMPADDING", (0, 0), (0, 0), 6)]))
    story += [info, Spacer(1, 5 * mm)]

    ageing = stmt.get("ageing")
    if ageing:
        cells = _ageing_cells(ageing)
        labels = [c[0] for c in cells]
        amounts = [_money(c[1]) for c in cells]
        age_tbl = Table([labels, amounts], colWidths=[width / len(labels)] * len(labels))
        age_tbl.setStyle(TableStyle([
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica"), ("FONTSIZE", (0, 0), (-1, 0), 7),
            ("TEXTCOLOR", (0, 0), (-1, 0), grey),
            ("FONTNAME", (0, 1), (-1, 1), "Helvetica-Bold"), ("FONTSIZE", (0, 1), (-1, 1), 8.5),
            ("ALIGN", (0, 0), (-1, -1), "CENTER"),
            ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#D1D5DB")),
            ("INNERGRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#E5E7EB")),
            ("BACKGROUND", (-1, 0), (-1, -1), colors.HexColor("#" + BRAND_LIGHT)),
        ]))
        story += [Paragraph(f"Ageing as at {_d(stmt.get('as_of'))}", small), Spacer(1, 1.5 * mm), age_tbl, Spacer(1, 5 * mm)]

    if is_open:
        header = ["Date", "Number", "Type", "Reference", "Due", "Days overdue", "Outstanding"]
        col_w = [0.11, 0.18, 0.14, 0.22, 0.11, 0.11, 0.13]
    else:
        header = ["Date", "Number", "Type", "Reference", "Due", "Debit", "Credit", "Balance"]
        col_w = [0.10, 0.17, 0.13, 0.18, 0.10, 0.10, 0.10, 0.12]
    cell = ParagraphStyle("cell", parent=base, fontSize=7.5, leading=9.5)
    data = [header]
    if not is_open and stmt.get("date_from"):
        data.append([_d(stmt["date_from"]), "", "", Paragraph("<i>Opening balance</i>", cell), "", "", "",
                     _money((stmt.get("summary") or {}).get("opening_balance", 0))])
    overdue_rows = []
    for r in stmt.get("rows", []):
        due = _d(r.get("due_date")) if r.get("kind") in ("invoice", "credit_note") else ""
        line = [_d(r.get("date")), Paragraph(r.get("number") or "", cell), r.get("type_label"),
                Paragraph((r.get("reference") or "")[:60], cell), due]
        if is_open:
            line += [str(r.get("days_overdue") or 0) if (r.get("residual") or 0) > 0 else "", _money(r.get("residual"))]
            if (r.get("days_overdue") or 0) > 0 and (r.get("residual") or 0) > 0:
                overdue_rows.append(len(data))
        else:
            line += [_money(r["debit"]) if r.get("debit") else "", _money(r["credit"]) if r.get("credit") else "",
                     _money(r.get("running_balance"))]
        data.append(line)
    if len(data) == 1:
        data.append(["", "", "", Paragraph("No transactions in this period.", cell), "", "", ""] + ([] if is_open else [""]))

    tx = Table(data, colWidths=[width * w for w in col_w], repeatRows=1)
    ts = [
        ("BACKGROUND", (0, 0), (-1, 0), brand),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 7.5),
        ("ALIGN", (5, 0), (-1, -1), "RIGHT"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LINEBELOW", (0, 1), (-1, -1), 0.25, colors.HexColor("#E5E7EB")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F9FAFB")]),
        ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]
    for i in overdue_rows:
        ts.append(("TEXTCOLOR", (5, i), (5, i), colors.HexColor("#" + RED)))
    tx.setStyle(TableStyle(ts))
    story += [tx, Spacer(1, 6 * mm)]

    if stmt.get("truncated"):
        story.append(Paragraph("This statement is too long to show in full. Please request a shorter period.",
                               ParagraphStyle("warn", parent=base, textColor=colors.HexColor("#" + RED))))
        story.append(Spacer(1, 3 * mm))

    story.append(Paragraph(_balance_note(stmt), bold))
    bank = company.get("bank") or {}
    if bank.get("account"):
        story.append(Spacer(1, 2 * mm))
        story.append(Paragraph(
            f"<b>Banking details</b><br/>{bank['account']}"
            + (f"<br/>Account holder: {bank['holder']}" if bank.get("holder") else ""), base))
    story.append(Spacer(1, 3 * mm))
    story.append(Paragraph(
        "If anything on this statement does not match your records, please contact us so we can resolve it.",
        small))

    generated = _d(stmt.get("generated_at"))
    co_name = company.get("name") or ""

    def _footer(canvas, doc_):
        canvas.saveState()
        canvas.setFont("Helvetica", 7)
        canvas.setFillColor(grey)
        canvas.drawString(14 * mm, 9 * mm, f"{co_name}  |  Statement generated {generated}")
        canvas.drawRightString(A4[0] - 14 * mm, 9 * mm, f"Page {doc_.page}")
        canvas.restoreState()

    doc.build(story, onFirstPage=_footer, onLaterPages=_footer)
    return buf.getvalue()
